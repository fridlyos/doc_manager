"""Extractor unit tests: PDF page numbers, text/CSV sections, error isolation."""

from __future__ import annotations

from pathlib import Path

import pymupdf
import pytest

from doc_manager.extraction import (
    ExtractionError,
    ExtractionErrorCode,
    get_extractor,
)
from doc_manager.extraction.csv_ import CsvExtractor
from doc_manager.extraction.docx_ import DocxExtractor
from doc_manager.extraction.pdf import PdfExtractor, PdfOcrOptions
from doc_manager.extraction.text import TextExtractor


def _make_pdf(path: Path, pages: list[str], *, password: str | None = None) -> None:
    doc = pymupdf.open()
    for body in pages:
        page = doc.new_page()
        if body:
            page.insert_text((72, 72), body)
    if password is not None:
        doc.save(
            path,
            encryption=pymupdf.PDF_ENCRYPT_AES_256,
            owner_pw=password,
            user_pw=password,
        )
    else:
        doc.save(path)
    doc.close()


def test_pdf_preserves_one_based_page_numbers(tmp_path: Path) -> None:
    pdf = tmp_path / "doc.pdf"
    _make_pdf(pdf, ["Alpha page one", "Bravo page two", "Charlie page three"])

    result = PdfExtractor().extract(pdf)

    assert result.page_count == 3
    assert [p.page_number for p in result.pages] == [1, 2, 3]
    assert [p.index for p in result.pages] == [0, 1, 2]
    assert "Alpha" in result.pages[0].text
    assert "Charlie" in result.pages[2].text
    assert result.metadata["page_count"] == 3


def test_pdf_encrypted_raises_specific_code(tmp_path: Path) -> None:
    pdf = tmp_path / "locked.pdf"
    _make_pdf(pdf, ["secret"], password="hunter2")
    with pytest.raises(ExtractionError) as exc:
        PdfExtractor().extract(pdf)
    assert exc.value.code is ExtractionErrorCode.encrypted


def test_pdf_image_only_reports_no_extractable_text(tmp_path: Path) -> None:
    pdf = tmp_path / "scanned.pdf"
    _make_pdf(pdf, ["", ""])  # pages with no text layer
    with pytest.raises(ExtractionError) as exc:
        PdfExtractor().extract(pdf)
    assert exc.value.code is ExtractionErrorCode.no_extractable_text


def test_pdf_ocr_recovers_image_only_pages(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    # OCR engine stubbed (no Tesseract needed): empty pages get OCR'd text instead
    # of failing with no_extractable_text.
    monkeypatch.setattr(PdfExtractor, "_ocr_page", lambda self, page: "scanned words here")
    pdf = tmp_path / "scanned.pdf"
    _make_pdf(pdf, ["", ""])  # no text layer
    result = PdfExtractor(PdfOcrOptions()).extract(pdf)
    assert result.page_count == 2
    assert all("scanned words here" in p.text for p in result.pages)
    assert result.metadata["ocr"] is True
    assert result.extractor_version == "pdf-1-ocr"


def test_pdf_ocr_skips_pages_with_text(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    # Text-layer pages must not invoke OCR even when OCR is enabled.
    def _boom(self: PdfExtractor, page: object) -> str:
        raise AssertionError("OCR must not run on a page with selectable text")

    monkeypatch.setattr(PdfExtractor, "_ocr_page", _boom)
    pdf = tmp_path / "text.pdf"
    _make_pdf(pdf, ["Alpha has real text"])
    result = PdfExtractor(PdfOcrOptions()).extract(pdf)
    assert "Alpha" in result.pages[0].text


def test_pdf_ocr_engine_failure_reports_ocr_failed(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    def _raise(*args: object, **kwargs: object) -> object:
        raise RuntimeError("tesseract not found")

    monkeypatch.setattr(pymupdf.Page, "get_textpage_ocr", _raise, raising=False)
    pdf = tmp_path / "scanned.pdf"
    _make_pdf(pdf, [""])
    with pytest.raises(ExtractionError) as exc:
        PdfExtractor(PdfOcrOptions()).extract(pdf)
    assert exc.value.code is ExtractionErrorCode.ocr_failed


def test_pdf_empty_file(tmp_path: Path) -> None:
    pdf = tmp_path / "empty.pdf"
    pdf.write_bytes(b"")
    with pytest.raises(ExtractionError) as exc:
        PdfExtractor().extract(pdf)
    assert exc.value.code is ExtractionErrorCode.empty_file


def test_pdf_malformed(tmp_path: Path) -> None:
    pdf = tmp_path / "junk.pdf"
    pdf.write_bytes(b"%PDF-1.4 not really a pdf at all")
    with pytest.raises(ExtractionError) as exc:
        PdfExtractor().extract(pdf)
    assert exc.value.code is ExtractionErrorCode.malformed


def test_text_splits_into_sections(tmp_path: Path) -> None:
    txt = tmp_path / "notes.txt"
    txt.write_text("First para line.\n\nSecond para here.\n\n\nThird para.")
    result = TextExtractor().extract(txt)
    assert result.page_count == 3
    assert all(p.page_number is None for p in result.pages)
    assert result.pages[0].text == "First para line."
    assert result.pages[2].text == "Third para."


def test_text_empty_file(tmp_path: Path) -> None:
    txt = tmp_path / "blank.txt"
    txt.write_text("   \n\n  ")
    with pytest.raises(ExtractionError) as exc:
        TextExtractor().extract(txt)
    assert exc.value.code is ExtractionErrorCode.empty_file


def test_docx_paragraphs_and_tables(tmp_path: Path) -> None:
    import docx

    doc = docx.Document()
    doc.add_paragraph("First Word paragraph.")
    doc.add_paragraph("")
    doc.add_paragraph("Second section.")
    table = doc.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Key"
    table.rows[0].cells[1].text = "Value"
    path = tmp_path / "memo.docx"
    doc.save(str(path))

    result = DocxExtractor().extract(path)
    assert result.extractor_name == "docx"
    assert all(p.page_number is None for p in result.pages)
    text = "\n".join(p.text for p in result.pages)
    assert "First Word paragraph." in text
    assert "Second section." in text
    assert "Key\tValue" in text


def test_docx_not_a_package(tmp_path: Path) -> None:
    bogus = tmp_path / "fake.docx"
    bogus.write_text("this is not a zip-based docx")
    with pytest.raises(ExtractionError) as exc:
        DocxExtractor().extract(bogus)
    assert exc.value.code is ExtractionErrorCode.malformed


def test_csv_repeats_header_and_is_row_aware(tmp_path: Path) -> None:
    csv_path = tmp_path / "data.csv"
    csv_path.write_text("name,role\nAda,engineer\nGrace,admiral\n")
    result = CsvExtractor().extract(csv_path)
    assert result.metadata["row_count"] == 2
    assert result.metadata["column_count"] == 2
    body = result.pages[0].text
    assert "name | role" in body  # header line
    assert "name: Ada | role: engineer" in body  # row-aware key: value


def test_csv_batches_rows_and_repeats_header(tmp_path: Path) -> None:
    rows = "\n".join(f"r{i},v{i}" for i in range(250))
    (tmp_path / "big.csv").write_text(f"a,b\n{rows}\n")
    result = CsvExtractor().extract(tmp_path / "big.csv")
    # 250 rows / 100 per section -> 3 sections, each starting with the header.
    assert result.page_count == 3
    assert all(page.text.startswith("a | b") for page in result.pages)


def test_registry_dispatch_and_unsupported() -> None:
    assert isinstance(get_extractor("pdf"), PdfExtractor)
    assert isinstance(get_extractor(".MD"), TextExtractor)
    assert isinstance(get_extractor("csv"), CsvExtractor)
    assert isinstance(get_extractor("docx"), DocxExtractor)
    # log files are no longer indexed; genuinely unsupported types return None.
    assert get_extractor("log") is None
    assert get_extractor("xlsx") is None


def test_registry_builds_ocr_pdf_when_enabled() -> None:
    from doc_manager.core.config import Settings

    off = get_extractor("pdf", settings=Settings(ocr_enabled=False))
    assert isinstance(off, PdfExtractor) and off.version == "pdf-1"

    on = get_extractor("pdf", settings=Settings(ocr_enabled=True, ocr_language="deu"))
    assert isinstance(on, PdfExtractor) and on.version == "pdf-1-ocr"
    assert on._ocr == PdfOcrOptions(language="deu")  # type: ignore[attr-defined]
    # OCR only affects PDFs; other types are unchanged.
    assert isinstance(get_extractor("txt", settings=Settings(ocr_enabled=True)), TextExtractor)
