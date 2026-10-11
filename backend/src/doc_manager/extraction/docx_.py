"""Word (.docx) extractor using python-docx.

Reads paragraph and table text from a modern Word document. ``.docx`` has no
fixed pages, so ``page_number`` is ``None``; text is split into synthetic
sections on blank-line boundaries, matching the plain-text extractor.
"""

from __future__ import annotations

import re
from pathlib import Path

from docx import Document as DocxDocument
from docx.opc.exceptions import PackageNotFoundError

from doc_manager.extraction.base import ExtractedDocument, ExtractedPage
from doc_manager.extraction.errors import ExtractionError, ExtractionErrorCode

_SECTION_SPLIT = re.compile(r"\n[ \t]*\n")


class DocxExtractor:
    name = "docx"
    version = "docx-1"

    def extract(self, path: Path) -> ExtractedDocument:
        try:
            doc = DocxDocument(str(path))
        except PackageNotFoundError as exc:
            # Not a valid .docx package (corrupt, or a legacy .doc renamed).
            raise ExtractionError(
                ExtractionErrorCode.malformed, "not a valid .docx document."
            ) from exc
        except Exception as exc:  # noqa: BLE001 - normalize python-docx errors
            raise ExtractionError(
                ExtractionErrorCode.malformed, f"could not read .docx: {type(exc).__name__}"
            ) from exc

        blocks: list[str] = [p.text for p in doc.paragraphs]
        for table in doc.tables:
            for row in table.rows:
                blocks.append("\t".join(cell.text for cell in row.cells))
        content = "\n".join(b for b in blocks).strip()
        if not content:
            raise ExtractionError(ExtractionErrorCode.no_extractable_text, "document has no text.")

        sections = [block.strip() for block in _SECTION_SPLIT.split(content)]
        sections = [block for block in sections if block]
        pages = [
            ExtractedPage(index=i, page_number=None, text=block) for i, block in enumerate(sections)
        ]
        return ExtractedDocument(
            extractor_name=self.name,
            extractor_version=self.version,
            pages=pages,
            metadata={"section_count": len(pages)},
        )
